import io
import json

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.ingestion.extraction import (
    DocumentExtractionError,
    extract_document_text,
    is_supported_document,
)


@pytest.mark.parametrize(
    ("filename", "content", "expected"),
    [
        ("report.txt", b"Sustainability priority", "Sustainability"),
        ("report.md", b"# Priority\nSustainability", "Priority"),
        ("metrics.csv", b"region,revenue\nEast,125000", "125000"),
        ("payload.json", json.dumps({"region": "East"}).encode(), "East"),
        ("page.html", b"<h1>Report</h1><p>Priority</p>", "Priority"),
    ],
)
def test_extracts_supported_text_formats(filename: str, content: bytes, expected: str) -> None:
    assert expected in extract_document_text(filename=filename, content=content)


def test_extracts_docx_and_xlsx() -> None:
    docx = Document()
    docx.add_paragraph("Sustainability priority")
    docx_buffer = io.BytesIO()
    docx.save(docx_buffer)

    workbook = Workbook()
    workbook.active.append(["region", "revenue"])
    workbook.active.append(["East", 125000])
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    assert "Sustainability" in extract_document_text(
        filename="report.docx", content=docx_buffer.getvalue()
    )
    extracted_xlsx = extract_document_text(
        filename="metrics.xlsx", content=xlsx_buffer.getvalue()
    )
    assert "125000" in extracted_xlsx


def test_rejects_encrypted_pdf_and_unsupported_format() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("secret")
    pdf = io.BytesIO()
    writer.write(pdf)

    with pytest.raises(DocumentExtractionError, match="Password-protected"):
        extract_document_text(filename="locked.pdf", content=pdf.getvalue())
    assert not is_supported_document("legacy.doc")
